"""
Scanner クラス群（html_builder.py から移植）

FolderScanner / PDFScanner / TextScanner / PlanScanner
+ CSV/Excel/Vega-Lite 変換ユーティリティ
"""

from __future__ import annotations

import csv
import json
import os
import re
import shutil
import sys
from pathlib import Path

from .content_bundle import (
    ContentBundle, ImageEntry, TableEntry, ChartEntry,
    DEFAULT_META, IMAGE_EXTENSIONS, TABLE_EXTENSIONS,
    CHART_EXTENSIONS, TEXT_EXTENSIONS, PDF_EXTENSIONS,
    classify_image_type, esc,
)

# オプション依存（インストールされていなければ機能を省略）
try:
    import fitz  # PyMuPDF
    HAS_PYMUPDF = True
except ImportError:
    HAS_PYMUPDF = False

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    from PIL import Image
    HAS_PILLOW = True
except ImportError:
    HAS_PILLOW = False

# 同パッケージの補助モジュール
SKILL_DIR = Path(__file__).parent
from .section_splitter import split_sections, split_text, get_ordered_sections, section_to_label


# ──────────────────────────────────────────────────
# ユーティリティ
# ──────────────────────────────────────────────────

def _detect_encoding(filepath: Path) -> str:
    """ファイルのエンコーディングを自動判定する"""
    for enc in ("utf-8-sig", "utf-8", "shift_jis", "cp932", "latin-1"):
        try:
            filepath.read_text(encoding=enc)
            return enc
        except (UnicodeDecodeError, LookupError):
            continue
    return "utf-8"


def _csv_to_html(path: Path) -> str:
    """CSV ファイルを HTML テーブルに変換する"""
    enc = _detect_encoding(path)
    rows = []
    try:
        with open(path, encoding=enc, errors="replace", newline="") as f:
            reader = csv.reader(f)
            rows = list(reader)
    except OSError:
        return ""

    if not rows:
        return ""

    header, body_rows = rows[0], rows[1:]
    th = "".join(f"<th>{esc(c)}</th>" for c in header)
    trs = ""
    for row in body_rows:
        tds = "".join(f"<td>{esc(c)}</td>" for c in row)
        trs += f"<tr>{tds}</tr>\n"
    return f"<table>\n<thead><tr>{th}</tr></thead>\n<tbody>\n{trs}</tbody>\n</table>"


def _xlsx_to_html(path: Path) -> str:
    """Excel ファイルを HTML テーブルに変換する（openpyxl 必要）"""
    if not HAS_OPENPYXL:
        print(f"警告: openpyxl がインストールされていません。{path.name} をスキップします。")
        print("  pip install openpyxl")
        return ""
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
    except Exception as e:
        print(f"警告: {path.name} の読み込みに失敗しました: {e}")
        return ""

    if not rows:
        return ""

    header, body_rows = rows[0], rows[1:]
    th = "".join(f"<th>{esc(str(c) if c is not None else '')}</th>" for c in header)
    trs = ""
    for row in body_rows:
        tds = "".join(f"<td>{esc(str(c) if c is not None else '')}</td>" for c in row)
        trs += f"<tr>{tds}</tr>\n"
    return f"<table>\n<thead><tr>{th}</tr></thead>\n<tbody>\n{trs}</tbody>\n</table>"


def _vegalite_to_svg(spec: dict, output_path: str) -> bool:
    """Vega-Lite JSON → SVG ファイルを生成する（vl-convert-python 必要）"""
    try:
        import vl_convert as vlc
        svg = vlc.vegalite_to_svg(vl_spec=spec)
        Path(output_path).write_text(svg, encoding="utf-8")
        return True
    except ImportError:
        print("警告: vl-convert-python がインストールされていません。グラフをスキップします。")
        print("  pip install vl-convert-python")
        return False
    except Exception as e:
        print(f"警告: グラフ生成に失敗しました: {e}")
        return False


# ──────────────────────────────────────────────────
# スキャナ
# ──────────────────────────────────────────────────

class FolderScanner:
    """
    フォルダをスキャンして ContentBundle を構築する。
    PDF / テキスト / 画像 / 表 / グラフJSON の複合入力に対応。
    """

    def __init__(self, folder: str | Path, output_dir: str | Path | None = None):
        self.folder = Path(folder)
        self.output_dir = Path(output_dir) if output_dir else self.folder / "output"

    def scan(self) -> ContentBundle:
        bundle = ContentBundle()
        bundle.output_dir = str(self.output_dir)

        # meta.json を先に読む
        meta = self._load_meta()
        bundle.title             = meta.get("title",             DEFAULT_META["title"])
        bundle.authors           = meta.get("authors",           DEFAULT_META["authors"])
        bundle.affiliation       = meta.get("affiliation",       DEFAULT_META["affiliation"])
        bundle.subtitle          = meta.get("subtitle",          DEFAULT_META["subtitle"])
        bundle.date              = meta.get("date",              DEFAULT_META["date"])
        bundle.presentation_type = meta.get("presentation_type", DEFAULT_META["presentation_type"])
        bundle.output_type       = meta.get("output_type",       DEFAULT_META["output_type"])
        bundle.theme             = meta.get("theme",             DEFAULT_META["theme"])
        bundle.language          = meta.get("language",          DEFAULT_META["language"])

        # 出力ディレクトリ準備
        images_dir = self.output_dir / "images"
        images_dir.mkdir(parents=True, exist_ok=True)

        # ファイルを種別ごとに仕分け
        text_files:  list[Path] = []
        image_files: list[Path] = []
        table_files: list[Path] = []
        chart_files: list[Path] = []
        pdf_files:   list[Path] = []

        for f in sorted(self.folder.iterdir()):
            if f.name == "meta.json":
                continue
            ext = f.suffix.lower()
            if ext in TEXT_EXTENSIONS:
                text_files.append(f)
            elif ext in IMAGE_EXTENSIONS:
                image_files.append(f)
            elif ext in TABLE_EXTENSIONS:
                table_files.append(f)
            elif ext in CHART_EXTENSIONS:
                chart_files.append(f)
            elif ext in PDF_EXTENSIONS:
                pdf_files.append(f)

        # テキスト処理
        bundle.sections = self._process_texts(text_files)

        # 画像処理
        bundle.images = self._process_images(image_files, images_dir)

        # 表処理
        bundle.tables = self._process_tables(table_files)

        # グラフ JSON 処理
        bundle.charts = self._process_charts(chart_files, images_dir)

        # PDF 処理（テキスト不足時の補完）
        for pdf in pdf_files:
            pdf_bundle = PDFScanner(pdf, self.output_dir).scan()
            if not bundle.sections:
                bundle.sections = pdf_bundle.sections
            bundle.images.extend(pdf_bundle.images)

        return bundle

    # ── meta.json ──

    def _load_meta(self) -> dict:
        meta_path = self.folder / "meta.json"
        if meta_path.exists():
            try:
                with open(meta_path, encoding="utf-8") as f:
                    return json.load(f)
            except (json.JSONDecodeError, OSError) as e:
                print(f"警告: meta.json の読み込みに失敗しました: {e}")
        return {}

    # ── テキスト ──

    def _process_texts(self, text_files: list[Path]) -> dict[str, str]:
        """テキストファイル群をセクション辞書に変換する"""
        from .section_splitter import _match_section_keyword

        sections: dict[str, str] = {}
        if not text_files:
            return sections

        if len(text_files) == 1:
            result = split_sections(text_files[0])
            if result:
                return result
            enc = _detect_encoding(text_files[0])
            return {"background": text_files[0].read_text(encoding=enc, errors="replace")}

        for tf in text_files:
            stem = tf.stem.lower()
            detected_section = _match_section_keyword(stem)
            enc = _detect_encoding(tf)
            content = tf.read_text(encoding=enc, errors="replace").strip()
            if not content:
                continue

            if detected_section:
                if detected_section in sections:
                    sections[detected_section] += "\n\n" + content
                else:
                    sections[detected_section] = content
            else:
                sub = split_text(content)
                if len(sub) >= 2:
                    for k, v in sub.items():
                        if k in sections:
                            sections[k] += "\n\n" + v
                        else:
                            sections[k] = v
                else:
                    key = re.sub(r'[^a-z0-9_]', '_', stem)
                    sections[key] = content

        return sections

    # ── 画像 ──

    def _process_images(
        self, image_files: list[Path], images_dir: Path
    ) -> list[ImageEntry]:
        entries = []
        for img in image_files:
            dest = images_dir / img.name
            if img != dest:
                shutil.copy2(img, dest)

            width, height, aspect = self._get_image_size(dest)
            entry = ImageEntry(
                path=str(dest).replace("\\", "/"),
                rel_path=f"images/{img.name}",
                filename=img.name,
                caption=self._infer_caption(img.stem),
                width=width,
                height=height,
                aspect_ratio=aspect,
                image_type=classify_image_type(aspect, img.stem),
            )
            entries.append(entry)
        return entries

    def _get_image_size(self, path: Path) -> tuple[int, int, float]:
        """画像の幅・高さ・アスペクト比を返す"""
        if HAS_PILLOW:
            try:
                with Image.open(path) as img:
                    w, h = img.size
                    return w, h, round(w / h, 3) if h > 0 else 1.0
            except Exception:
                pass
        return 0, 0, 1.0

    def _infer_caption(self, stem: str) -> str:
        """ファイル名から推定キャプションを生成する"""
        m = re.match(r'(fig|figure|table|chart|graph|photo)_?(\d+)', stem, re.IGNORECASE)
        if m:
            kind = m.group(1).capitalize()
            num  = m.group(2)
            label = "Figure" if kind.lower() in ("fig", "figure", "chart", "graph", "photo") else "Table"
            return f"{label} {num}."
        return stem.replace("_", " ").replace("-", " ")

    # ── 表データ ──

    def _process_tables(self, table_files: list[Path]) -> list[TableEntry]:
        entries = []
        for tf in table_files:
            ext = tf.suffix.lower()
            if ext == ".csv":
                html = _csv_to_html(tf)
            elif ext in (".xlsx", ".xls"):
                html = _xlsx_to_html(tf)
            else:
                continue
            if html:
                entries.append(TableEntry(
                    html=html,
                    caption=self._infer_caption(tf.stem),
                    source_file=tf.name,
                ))
        return entries

    # ── グラフ JSON ──

    def _process_charts(
        self, chart_files: list[Path], images_dir: Path
    ) -> list[ChartEntry]:
        entries = []
        for cf in chart_files:
            try:
                with open(cf, encoding="utf-8") as f:
                    spec = json.load(f)
                if "$schema" not in spec:
                    continue
                svg_name = cf.stem + ".svg"
                svg_path = images_dir / svg_name
                success = _vegalite_to_svg(spec, str(svg_path))
                if success:
                    entries.append(ChartEntry(
                        svg_path=str(svg_path).replace("\\", "/"),
                        rel_path=f"images/{svg_name}",
                        caption=self._infer_caption(cf.stem),
                        source_json=cf.name,
                    ))
            except (json.JSONDecodeError, OSError):
                continue
        return entries


class PDFScanner:
    """
    PDF ファイルをスキャンして ContentBundle を構築する。
    PyMuPDF が必要。
    """

    def __init__(self, pdf_path: str | Path, output_dir: str | Path | None = None):
        self.pdf_path = Path(pdf_path)
        self.output_dir = Path(output_dir) if output_dir else self.pdf_path.parent / self.pdf_path.stem / "output"

    def scan(self) -> ContentBundle:
        bundle = ContentBundle()
        bundle.output_dir = str(self.output_dir)

        if not HAS_PYMUPDF:
            print("警告: PyMuPDF がインストールされていません。PDF処理をスキップします。")
            print("  pip install pymupdf")
            return bundle

        images_dir = self.output_dir / "images"
        images_dir.mkdir(parents=True, exist_ok=True)

        doc = fitz.open(str(self.pdf_path))

        # テキスト抽出・セクション分割
        full_text = "\n\n".join(
            page.get_text() for page in doc
        ).strip()
        if full_text:
            bundle.sections = split_text(full_text)

        # 1ページ目からタイトル・著者を推定
        first_page_text = doc[0].get_text().strip() if len(doc) > 0 else ""
        if first_page_text:
            lines = [l.strip() for l in first_page_text.split("\n") if l.strip()]
            if lines:
                bundle.title = lines[0]
            if len(lines) >= 2:
                bundle.authors = lines[1]
            if len(lines) >= 3:
                bundle.affiliation = lines[2]

        # 図版抽出
        extract_script = SKILL_DIR / "extract_images.py"
        if extract_script.exists():
            import subprocess
            env = os.environ.copy()
            env["PYTHONIOENCODING"] = "utf-8"
            subprocess.run(
                [sys.executable, str(extract_script),
                 str(self.pdf_path), str(images_dir)],
                capture_output=True, text=True, encoding="utf-8",
                errors="replace", env=env
            )
            # images_index.json を読み込み
            index_path = images_dir / "images_index.json"
            if index_path.exists():
                with open(index_path, encoding="utf-8") as f:
                    index = json.load(f)
                for item in index:
                    bundle.images.append(ImageEntry(
                        path=item.get("path", ""),
                        rel_path=f"images/{item.get('filename', '')}",
                        filename=item.get("filename", ""),
                        caption=item.get("caption", ""),
                        width=item.get("width", 0),
                        height=item.get("height", 0),
                        aspect_ratio=item.get("aspect_ratio", 1.0),
                        image_type=item.get("type", "figure"),
                    ))

        doc.close()
        return bundle


class TextScanner:
    """
    単一テキストファイルをスキャンして ContentBundle を構築する。
    section_splitter でセクション自動分割する。
    """

    def __init__(self, text_path: str | Path, output_dir: str | Path | None = None):
        self.text_path = Path(text_path)
        self.output_dir = Path(output_dir) if output_dir else self.text_path.parent / self.text_path.stem / "output"

    def scan(self) -> ContentBundle:
        bundle = ContentBundle()
        bundle.output_dir = str(self.output_dir)
        bundle.sections = split_sections(self.text_path)
        return bundle


class PlanScanner:
    """
    slide_plan.json を読み込んで ContentBundle を構築する。
    """

    def __init__(self, plan_path: str | Path, output_dir: str | Path | None = None):
        self.plan_path = Path(plan_path)
        self.output_dir = Path(output_dir) if output_dir else self.plan_path.parent

    def scan(self) -> ContentBundle:
        with open(self.plan_path, encoding="utf-8") as f:
            plan = json.load(f)

        bundle = ContentBundle()
        bundle.output_dir = str(self.output_dir)

        meta = plan.get("meta", {})
        bundle.title             = meta.get("title",             DEFAULT_META["title"])
        bundle.authors           = meta.get("authors",           DEFAULT_META["authors"])
        bundle.affiliation       = meta.get("affiliation",       DEFAULT_META["affiliation"])
        bundle.subtitle          = meta.get("subtitle",          DEFAULT_META["subtitle"])
        bundle.date              = meta.get("date",              DEFAULT_META["date"])
        bundle.theme             = meta.get("theme",             DEFAULT_META["theme"])
        bundle.presentation_type = meta.get("presentation_type", DEFAULT_META["presentation_type"])
        bundle.language          = meta.get("language",          DEFAULT_META["language"])

        bundle._plan_slides = plan.get("slides", [])

        return bundle
